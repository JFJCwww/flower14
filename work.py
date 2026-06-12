# -*- coding: utf-8 -*-
import sys
import os
import warnings
from datetime import datetime
from PyQt5 import QtCore, QtGui, QtWidgets
from ultralytics import YOLO

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    print("请安装 openpyxl: pip install openpyxl")
    sys.exit(1)

try:
    import matplotlib
    matplotlib.use("Qt5Agg")
    from matplotlib.backends.backend_qt5agg import FigureCanvasQTAgg as FigureCanvas
    from matplotlib.figure import Figure
    import matplotlib.pyplot as plt
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei"]
    plt.rcParams["axes.unicode_minus"] = False
except ImportError:
    print("请安装 matplotlib: pip install matplotlib")
    sys.exit(1)

warnings.filterwarnings("ignore", category=DeprecationWarning)

# 14 种花的中英文映射
FLOWER_CN = {
    "astilbe": "落新妇",
    "bellflower": "风铃草",
    "black_eyed_susan": "黑眼菊",
    "calendula": "金盏花",
    "california_poppy": "金英花",
    "carnation": "康乃馨",
    "common_daisy": "雏菊",
    "coreopsis": "金鸡菊",
    "dandelion": "蒲公英",
    "iris": "鸢尾花",
    "rose": "玫瑰",
    "sunflower": "向日葵",
    "tulip": "郁金香",
    "water_lily": "睡莲",
}

MODEL_PATH = "runs/yolo_cls_best.pt"

# 全局样式表
STYLE_SHEET = """
QWidget {
    font-family: "Microsoft YaHei", "Segoe UI", sans-serif;
    background-color: #f5f7fa;
}

/* 标题栏 */
QLabel#label_title {
    background: qlineargradient(x1:0, y1:0, x2:1, y2:0,
        stop:0 #667eea, stop:1 #764ba2);
    color: white;
    font-size: 22px;
    font-weight: bold;
    padding: 16px;
    border-radius: 12px;
}

/* 图片显示区 */
QLabel#label_image {
    background-color: white;
    border: 2px solid #e0e6ed;
    border-radius: 14px;
    padding: 8px;
}

/* 识别结果卡片 */
QFrame#result_card {
    background-color: white;
    border-radius: 12px;
    border: 1px solid #e0e6ed;
}

QLabel#label_result_title {
    color: #667eea;
    font-size: 14px;
    font-weight: bold;
    padding: 0;
}

QLabel#label_result_flower {
    color: #2d3748;
    font-size: 28px;
    font-weight: bold;
    padding: 4px 0;
}

QLabel#label_result_conf {
    color: #68d391;
    font-size: 20px;
    font-weight: bold;
    padding: 2px 0;
}

/* 文件信息卡片 */
QFrame#file_card {
    background-color: white;
    border-radius: 12px;
    border: 1px solid #e0e6ed;
}

QLabel#label_file_title {
    color: #a0aec0;
    font-size: 13px;
    padding: 0;
}

QLabel#label_filename {
    color: #4a5568;
    font-size: 14px;
    padding: 2px 0;
}

QLabel#label_progress {
    color: #a0aec0;
    font-size: 13px;
    padding: 0;
}

/* 按钮 */
QPushButton {
    font-size: 15px;
    font-weight: bold;
    color: white;
    border: none;
    border-radius: 12px;
    padding: 14px 24px;
}

QPushButton#btn_start {
    background: qlineargradient(x1:0, y1:0, x2:1, y2:0,
        stop:0 #667eea, stop:1 #764ba2);
}

QPushButton#btn_start:hover {
    background: qlineargradient(x1:0, y1:0, x2:1, y2:0,
        stop:0 #5a6fd6, stop:1 #6a4192);
}

QPushButton#btn_start:pressed {
    background: qlineargradient(x1:0, y1:0, x2:1, y2:0,
        stop:0 #4e60c2, stop:1 #5e3782);
}

QPushButton#btn_stop {
    background: qlineargradient(x1:0, y1:0, x2:1, y2:0,
        stop:0 #fc8181, stop:1 #f687b3);
}

QPushButton#btn_stop:hover {
    background: qlineargradient(x1:0, y1:0, x2:1, y2:0,
        stop:0 #e06060, stop:1 #e07098);
}

QPushButton#btn_stop:pressed {
    background: qlineargradient(x1:0, y1:0, x2:1, y2:0,
        stop:0 #c05050, stop:1 #c06080);
}

QPushButton#btn_prev, QPushButton#btn_next {
    background-color: #edf2f7;
    color: #4a5568;
    font-size: 20px;
    padding: 10px 18px;
    border: 1px solid #e2e8f0;
}

QPushButton#btn_prev:hover, QPushButton#btn_next:hover {
    background-color: #e2e8f0;
}

/* 花朵统计面板 */
QFrame#stats_card {
    background-color: white;
    border-radius: 12px;
    border: 1px solid #e0e6ed;
}

QLabel#label_stats_title {
    color: #667eea;
    font-size: 14px;
    font-weight: bold;
    padding: 2px 0;
}

QScrollArea#stats_scroll {
    background-color: transparent;
    border: none;
}

QFrame#stat_row {
    background-color: transparent;
    border-radius: 6px;
    padding: 2px 0;
}

QLabel#stat_name {
    color: #4a5568;
    font-size: 13px;
}

QLabel#stat_count {
    color: #667eea;
    font-size: 13px;
    font-weight: bold;
}

QLabel#stat_conf {
    color: #68d391;
    font-size: 13px;
}

/* 进度条 */
QProgressBar {
    border: none;
    border-radius: 5px;
    background-color: #e2e8f0;
    height: 10px;
    text-align: center;
}

QProgressBar::chunk {
    background: qlineargradient(x1:0, y1:0, x2:1, y2:0,
        stop:0 #667eea, stop:1 #764ba2);
    border-radius: 5px;
}

/* 底部信息 */
QLabel#label_footer {
    color: #a0aec0;
    font-size: 12px;
}
"""


class ImagePlayer(QtWidgets.QWidget):
    def __init__(self):
        super().__init__()

        self.setObjectName("Form")
        self.setMinimumSize(960, 680)
        self.resize(1020, 720)
        self.setWindowTitle("花朵识别器")
        self.setStyleSheet(STYLE_SHEET)

        # 初始化变量
        self.image_list = []
        self.current_index = 0
        self.classified_count = 0
        self.flower_stats = {name: {"count": 0, "total_conf": 0.0} for name in FLOWER_CN}
        self.recognition_records = []
        self.timer = QtCore.QTimer(self)
        self.timer.setInterval(2000)

        # 加载花朵分类模型
        self.model = None
        if os.path.exists(MODEL_PATH):
            self.model = YOLO(MODEL_PATH)
            print(f"模型加载成功: {MODEL_PATH}")
        else:
            print(f"模型文件不存在: {MODEL_PATH}，请先训练模型")

        self.setup_ui()

        # 信号连接
        self.btn_start.clicked.connect(self.start_play)
        self.btn_stop.clicked.connect(self.stop_play)
        self.btn_prev.clicked.connect(self.prev_image)
        self.btn_next.clicked.connect(self.next_image)
        self.timer.timeout.connect(self.auto_next)
        self.btn_save.clicked.connect(self.save_to_excel)
        self.btn_import.clicked.connect(self.import_and_visualize)

        self.label_result_flower.setText("等待识别")
        self.label_result_conf.setText("--")
        self.label_filename.setText("请选择图片文件夹")
        self.label_progress.setText("0 / 0")

    def setup_ui(self):
        main_layout = QtWidgets.QVBoxLayout(self)
        main_layout.setContentsMargins(20, 16, 20, 16)
        main_layout.setSpacing(12)

        # === 标题 ===
        self.label_title = QtWidgets.QLabel("花朵识别器  Flower Classifier")
        self.label_title.setObjectName("label_title")
        self.label_title.setAlignment(QtCore.Qt.AlignCenter)
        main_layout.addWidget(self.label_title)

        # === 中间区域：图片 + 右侧面板 ===
        mid_layout = QtWidgets.QHBoxLayout()
        mid_layout.setSpacing(16)

        # -- 左侧：图片显示 --
        self.label_image = QtWidgets.QLabel("点击「开始识别」选择图片文件夹")
        self.label_image.setObjectName("label_image")
        self.label_image.setAlignment(QtCore.Qt.AlignCenter)
        self.label_image.setMinimumSize(480, 400)
        self.label_image.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Expanding)
        mid_layout.addWidget(self.label_image, stretch=3)

        # -- 右侧面板 --
        right_layout = QtWidgets.QVBoxLayout()
        right_layout.setSpacing(12)
        right_layout.setContentsMargins(0, 0, 0, 0)

        # 识别结果卡片
        result_card = QtWidgets.QFrame()
        result_card.setObjectName("result_card")
        result_card.setFixedWidth(260)
        result_layout = QtWidgets.QVBoxLayout(result_card)
        result_layout.setContentsMargins(20, 16, 20, 16)
        result_layout.setSpacing(6)

        self.label_result_title = QtWidgets.QLabel("识别结果")
        self.label_result_title.setObjectName("label_result_title")
        self.label_result_title.setAlignment(QtCore.Qt.AlignCenter)
        result_layout.addWidget(self.label_result_title)

        self.label_result_flower = QtWidgets.QLabel("等待识别")
        self.label_result_flower.setObjectName("label_result_flower")
        self.label_result_flower.setAlignment(QtCore.Qt.AlignCenter)
        result_layout.addWidget(self.label_result_flower)

        self.label_result_conf = QtWidgets.QLabel("--")
        self.label_result_conf.setObjectName("label_result_conf")
        self.label_result_conf.setAlignment(QtCore.Qt.AlignCenter)
        result_layout.addWidget(self.label_result_conf)

        right_layout.addWidget(result_card)

        # 文件信息卡片
        file_card = QtWidgets.QFrame()
        file_card.setObjectName("file_card")
        file_card.setFixedWidth(260)
        file_layout = QtWidgets.QVBoxLayout(file_card)
        file_layout.setContentsMargins(20, 14, 20, 14)
        file_layout.setSpacing(6)

        self.label_file_title = QtWidgets.QLabel("当前文件")
        self.label_file_title.setObjectName("label_file_title")
        file_layout.addWidget(self.label_file_title)

        self.label_filename = QtWidgets.QLabel("请选择图片文件夹")
        self.label_filename.setObjectName("label_filename")
        self.label_filename.setWordWrap(True)
        file_layout.addWidget(self.label_filename)

        self.label_progress = QtWidgets.QLabel("0 / 0")
        self.label_progress.setObjectName("label_progress")
        self.label_progress.setAlignment(QtCore.Qt.AlignRight)
        file_layout.addWidget(self.label_progress)

        # 进度条
        self.progress_bar = QtWidgets.QProgressBar()
        self.progress_bar.setFixedHeight(10)
        self.progress_bar.setTextVisible(False)
        self.progress_bar.setValue(0)
        file_layout.addWidget(self.progress_bar)

        right_layout.addWidget(file_card)

        # 花朵统计卡片
        stats_card = QtWidgets.QFrame()
        stats_card.setObjectName("stats_card")
        stats_card.setFixedWidth(260)
        stats_outer = QtWidgets.QVBoxLayout(stats_card)
        stats_outer.setContentsMargins(14, 12, 14, 12)
        stats_outer.setSpacing(8)

        self.label_stats_title = QtWidgets.QLabel("各花统计（数量 / 平均置信度）")
        self.label_stats_title.setObjectName("label_stats_title")
        self.label_stats_title.setAlignment(QtCore.Qt.AlignCenter)
        stats_outer.addWidget(self.label_stats_title)

        scroll = QtWidgets.QScrollArea()
        scroll.setObjectName("stats_scroll")
        scroll.setWidgetResizable(True)
        scroll.setHorizontalScrollBarPolicy(QtCore.Qt.ScrollBarAlwaysOff)
        scroll.setVerticalScrollBarPolicy(QtCore.Qt.ScrollBarAsNeeded)

        scroll_widget = QtWidgets.QWidget()
        scroll_widget.setStyleSheet("background: transparent;")
        self.stats_layout = QtWidgets.QVBoxLayout(scroll_widget)
        self.stats_layout.setContentsMargins(0, 0, 0, 0)
        self.stats_layout.setSpacing(2)

        self.stat_labels = {}
        for en_name, cn_name in FLOWER_CN.items():
            row = QtWidgets.QFrame()
            row.setObjectName("stat_row")
            row.setFixedHeight(28)
            row_layout = QtWidgets.QHBoxLayout(row)
            row_layout.setContentsMargins(8, 0, 8, 0)
            row_layout.setSpacing(0)

            lbl_name = QtWidgets.QLabel(cn_name)
            lbl_name.setObjectName("stat_name")
            row_layout.addWidget(lbl_name)

            lbl_count = QtWidgets.QLabel("0")
            lbl_count.setObjectName("stat_count")
            lbl_count.setFixedWidth(40)
            lbl_count.setAlignment(QtCore.Qt.AlignCenter)
            row_layout.addWidget(lbl_count)

            lbl_conf = QtWidgets.QLabel("--")
            lbl_conf.setObjectName("stat_conf")
            lbl_conf.setFixedWidth(65)
            lbl_conf.setAlignment(QtCore.Qt.AlignRight | QtCore.Qt.AlignVCenter)
            row_layout.addWidget(lbl_conf)

            self.stats_layout.addWidget(row)
            self.stat_labels[en_name] = (lbl_count, lbl_conf)

        self.stats_layout.addStretch()
        scroll.setWidget(scroll_widget)
        stats_outer.addWidget(scroll)

        right_layout.addWidget(stats_card, stretch=1)

        mid_layout.addLayout(right_layout, stretch=1)
        main_layout.addLayout(mid_layout, stretch=1)

        # === 底部控制栏 ===
        ctrl_layout = QtWidgets.QHBoxLayout()
        ctrl_layout.setSpacing(12)

        self.btn_prev = QtWidgets.QPushButton("<")
        self.btn_prev.setObjectName("btn_prev")
        self.btn_prev.setFixedSize(50, 50)
        self.btn_prev.setEnabled(False)
        ctrl_layout.addWidget(self.btn_prev)

        self.btn_start = QtWidgets.QPushButton("开始识别")
        self.btn_start.setObjectName("btn_start")
        self.btn_start.setMinimumHeight(50)
        ctrl_layout.addWidget(self.btn_start)

        self.btn_stop = QtWidgets.QPushButton("暂停")
        self.btn_stop.setObjectName("btn_stop")
        self.btn_stop.setMinimumHeight(50)
        self.btn_stop.setEnabled(False)
        ctrl_layout.addWidget(self.btn_stop)

        self.btn_next = QtWidgets.QPushButton(">")
        self.btn_next.setObjectName("btn_next")
        self.btn_next.setFixedSize(50, 50)
        self.btn_next.setEnabled(False)
        ctrl_layout.addWidget(self.btn_next)

        self.btn_save = QtWidgets.QPushButton("保存结果")
        self.btn_save.setObjectName("btn_start")
        self.btn_save.setMinimumHeight(50)
        self.btn_save.setEnabled(False)
        ctrl_layout.addWidget(self.btn_save)

        self.btn_import = QtWidgets.QPushButton("导入可视化")
        self.btn_import.setObjectName("btn_start")
        self.btn_import.setMinimumHeight(50)
        ctrl_layout.addWidget(self.btn_import)

        main_layout.addLayout(ctrl_layout)

        # === 底部提示 ===
        self.label_footer = QtWidgets.QLabel("Powered by YOLOv8  |  14 类花朵分类")
        self.label_footer.setObjectName("label_footer")
        self.label_footer.setAlignment(QtCore.Qt.AlignCenter)
        main_layout.addWidget(self.label_footer)

    def start_play(self):
        folder_path = QtWidgets.QFileDialog.getExistingDirectory(
            self, "选择图片文件夹", "./"
        )
        if not folder_path:
            return

        self.image_list = []
        for filename in os.listdir(folder_path):
            if filename.lower().endswith(('.png', '.jpg', '.jpeg', '.bmp', '.gif', '.webp')):
                self.image_list.append(os.path.join(folder_path, filename))

        if not self.image_list:
            QtWidgets.QMessageBox.warning(self, "警告", "所选文件夹中没有找到图片文件！")
            return

        self.current_index = 0
        self.classified_count = 0
        self.flower_stats = {name: {"count": 0, "total_conf": 0.0} for name in FLOWER_CN}
        for lbl_count, lbl_conf in self.stat_labels.values():
            lbl_count.setText("0")
            lbl_conf.setText("--")
        self.progress_bar.setMaximum(len(self.image_list))
        self.show_current_image()
        self.timer.start()

        self.btn_stop.setEnabled(True)
        self.btn_prev.setEnabled(True)
        self.btn_next.setEnabled(True)
        self.btn_start.setText("重新选择")

    def stop_play(self):
        if self.timer.isActive():
            self.timer.stop()
            self.btn_stop.setText("继续")
        else:
            self.timer.start()
            self.btn_stop.setText("暂停")

    def auto_next(self):
        if self.current_index < len(self.image_list) - 1:
            self.current_index += 1
            self.show_current_image()
        else:
            self.timer.stop()
            self.btn_stop.setText("继续")
            self.label_result_flower.setText("播放完毕")
            self.label_result_conf.setText("")

    def prev_image(self):
        if not self.image_list:
            return
        if self.timer.isActive():
            self.timer.stop()
            self.btn_stop.setText("继续")
        self.current_index = max(0, self.current_index - 1)
        self.show_current_image()

    def next_image(self):
        if not self.image_list:
            return
        if self.timer.isActive():
            self.timer.stop()
            self.btn_stop.setText("继续")
        self.current_index = min(len(self.image_list) - 1, self.current_index + 1)
        self.show_current_image()

    def predict_flower(self, image_path):
        if self.model is None:
            return None, "模型未加载", ""
        try:
            results = self.model(image_path, verbose=False)
            r = results[0]
            cls_id = int(r.probs.top1)
            cls_name = r.names[cls_id]
            confidence = r.probs.top1conf.item()
            cn_name = FLOWER_CN.get(cls_name, cls_name)
            return cls_name, cn_name, f"{confidence:.1%}"
        except Exception as e:
            return None, "识别失败", str(e)

    def show_current_image(self):
        image_path = self.image_list[self.current_index]

        pixmap = QtGui.QPixmap(image_path)
        scaled_pixmap = pixmap.scaled(
            self.label_image.size(),
            QtCore.Qt.KeepAspectRatio,
            QtCore.Qt.SmoothTransformation
        )
        self.label_image.setPixmap(scaled_pixmap)

        self.label_filename.setText(os.path.basename(image_path))
        self.label_progress.setText(f"{self.current_index + 1} / {len(self.image_list)}")
        self.progress_bar.setValue(self.current_index + 1)

        en_name, cn_name, confidence = self.predict_flower(image_path)
        self.label_result_flower.setText(cn_name)
        if confidence:
            self.label_result_conf.setText(confidence)
            conf_val = float(confidence.strip("%")) / 100
            if en_name and en_name in self.flower_stats:
                self.flower_stats[en_name]["count"] += 1
                self.flower_stats[en_name]["total_conf"] += conf_val
                self._update_stat_row(en_name)
            self.classified_count += 1
            self.recognition_records.append({
                "文件名": os.path.basename(image_path),
                "文件路径": image_path,
                "英文名": en_name or "",
                "中文名": cn_name,
                "置信度": conf_val,
                "识别时间": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            })
            self.btn_save.setEnabled(True)
        else:
            self.label_result_conf.setText("--")

    def _update_stat_row(self, en_name):
        stats = self.flower_stats[en_name]
        lbl_count, lbl_conf = self.stat_labels[en_name]
        lbl_count.setText(str(stats["count"]))
        avg = stats["total_conf"] / stats["count"]
        lbl_conf.setText(f"{avg:.1%}")

    def save_to_excel(self):
        if not self.recognition_records:
            QtWidgets.QMessageBox.warning(self, "警告", "没有可保存的识别记录！")
            return

        save_path, _ = QtWidgets.QFileDialog.getSaveFileName(
            self, "保存识别结果", "识别结果.xlsx", "Excel 文件 (*.xlsx)"
        )
        if not save_path:
            return

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "识别记录"
            headers = ["文件名", "文件路径", "英文名", "中文名", "置信度", "识别时间"]
            ws.append(headers)
            for cell in ws[1]:
                cell.font = cell.font.copy(bold=True)

            for record in self.recognition_records:
                ws.append([
                    record["文件名"],
                    record["文件路径"],
                    record["英文名"],
                    record["中文名"],
                    record["置信度"],
                    record["识别时间"],
                ])

            # 汇总统计表
            ws2 = wb.create_sheet("统计汇总")
            ws2.append(["花朵名称", "英文名", "识别次数", "平均置信度"])
            for cell in ws2[1]:
                cell.font = cell.font.copy(bold=True)

            for en_name, cn_name in FLOWER_CN.items():
                stats = self.flower_stats[en_name]
                if stats["count"] > 0:
                    avg_conf = stats["total_conf"] / stats["count"]
                else:
                    avg_conf = 0
                ws2.append([cn_name, en_name, stats["count"], f"{avg_conf:.1%}"])

            wb.save(save_path)
            QtWidgets.QMessageBox.information(
                self, "保存成功", f"识别结果已保存到：\n{save_path}"
            )
        except Exception as e:
            QtWidgets.QMessageBox.critical(self, "保存失败", f"保存时发生错误：\n{e}")

    def import_and_visualize(self):
        file_path, _ = QtWidgets.QFileDialog.getOpenFileName(
            self, "选择 Excel 文件", "./", "Excel 文件 (*.xlsx)"
        )
        if not file_path:
            return

        try:
            wb = load_workbook(file_path)
            if "识别记录" not in wb.sheetnames:
                QtWidgets.QMessageBox.warning(self, "警告", "该 Excel 文件中没有找到「识别记录」工作表！")
                return

            ws = wb["识别记录"]
            records = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is None:
                    continue
                records.append({
                    "文件名": row[0],
                    "英文名": row[2] if len(row) > 2 else "",
                    "中文名": row[3] if len(row) > 3 else "",
                    "置信度": row[4] if len(row) > 4 else 0,
                })

            if not records:
                QtWidgets.QMessageBox.warning(self, "警告", "识别记录为空！")
                return

            self.viz_window = VisualizationWindow(records)
            self.viz_window.show()

        except Exception as e:
            QtWidgets.QMessageBox.critical(self, "导入失败", f"读取 Excel 时发生错误：\n{e}")


class VisualizationWindow(QtWidgets.QWidget):
    def __init__(self, records):
        super().__init__()
        self.records = records
        self.setWindowTitle("识别数据可视化")
        self.setMinimumSize(1000, 700)
        self.setStyleSheet(STYLE_SHEET)

        layout = QtWidgets.QVBoxLayout(self)
        layout.setContentsMargins(16, 12, 16, 12)
        layout.setSpacing(10)

        title = QtWidgets.QLabel("花朵识别数据可视化")
        title.setObjectName("label_title")
        title.setAlignment(QtCore.Qt.AlignCenter)
        layout.addWidget(title)

        # 图表选择
        btn_layout = QtWidgets.QHBoxLayout()
        btn_layout.setSpacing(10)

        self.btn_bar = QtWidgets.QPushButton("数量柱状图")
        self.btn_bar.setObjectName("btn_start")
        self.btn_bar.clicked.connect(self.plot_bar)
        btn_layout.addWidget(self.btn_bar)

        self.btn_pie = QtWidgets.QPushButton("占比饼图")
        self.btn_pie.setObjectName("btn_start")
        self.btn_pie.clicked.connect(self.plot_pie)
        btn_layout.addWidget(self.btn_pie)

        self.btn_conf = QtWidgets.QPushButton("置信度分布")
        self.btn_conf.setObjectName("btn_start")
        self.btn_conf.clicked.connect(self.plot_confidence)
        btn_layout.addWidget(self.btn_conf)

        self.btn_line = QtWidgets.QPushButton("置信度折线")
        self.btn_line.setObjectName("btn_start")
        self.btn_line.clicked.connect(self.plot_line)
        btn_layout.addWidget(self.btn_line)

        layout.addLayout(btn_layout)

        # matplotlib 画布
        self.figure = Figure(figsize=(10, 5), dpi=100)
        self.canvas = FigureCanvas(self.figure)
        layout.addWidget(self.canvas, stretch=1)

        # 数据表格
        self.table = QtWidgets.QTableWidget()
        self.table.setColumnCount(4)
        self.table.setHorizontalHeaderLabels(["文件名", "中文名", "英文名", "置信度"])
        self.table.horizontalHeader().setStretchLastSection(True)
        self.table.setEditTriggers(QtWidgets.QAbstractItemView.NoEditTriggers)
        self.table.setSelectionBehavior(QtWidgets.QAbstractItemView.SelectRows)
        self.table.setMaximumHeight(180)
        layout.addWidget(self.table)

        self._populate_table()
        self.plot_bar()

    def _populate_table(self):
        self.table.setRowCount(len(self.records))
        for i, rec in enumerate(self.records):
            self.table.setItem(i, 0, QtWidgets.QTableWidgetItem(str(rec["文件名"])))
            self.table.setItem(i, 1, QtWidgets.QTableWidgetItem(str(rec["中文名"])))
            self.table.setItem(i, 2, QtWidgets.QTableWidgetItem(str(rec["英文名"])))
            conf = rec["置信度"]
            if isinstance(conf, (int, float)):
                conf_text = f"{conf:.1%}"
            else:
                conf_text = str(conf)
            self.table.setItem(i, 3, QtWidgets.QTableWidgetItem(conf_text))

    def _aggregate(self):
        agg = {}
        for rec in self.records:
            name = rec["中文名"]
            if name not in agg:
                agg[name] = {"count": 0, "total_conf": 0.0}
            agg[name]["count"] += 1
            conf = rec["置信度"]
            if isinstance(conf, (int, float)):
                agg[name]["total_conf"] += conf
        return agg

    def plot_bar(self):
        self.figure.clear()
        ax = self.figure.add_subplot(111)
        agg = self._aggregate()
        names = list(agg.keys())
        counts = [agg[n]["count"] for n in names]
        colors = plt.cm.Set3([i / max(len(names), 1) for i in range(len(names))])

        bars = ax.bar(names, counts, color=colors, edgecolor="#333", linewidth=0.5)
        for bar, count in zip(bars, counts):
            ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + 0.3,
                    str(count), ha="center", va="bottom", fontsize=10, fontweight="bold")

        ax.set_title("各花朵识别数量", fontsize=14, fontweight="bold")
        ax.set_ylabel("识别次数")
        ax.set_xlabel("花朵名称")
        self.figure.tight_layout()
        self.canvas.draw()

    def plot_pie(self):
        self.figure.clear()
        ax = self.figure.add_subplot(111)
        agg = self._aggregate()
        names = list(agg.keys())
        counts = [agg[n]["count"] for n in names]
        colors = plt.cm.Set3([i / max(len(names), 1) for i in range(len(names))])

        wedges, texts, autotexts = ax.pie(
            counts, labels=names, autopct="%1.1f%%",
            colors=colors, startangle=90, pctdistance=0.85
        )
        for t in autotexts:
            t.set_fontsize(9)
        ax.set_title("花朵识别占比", fontsize=14, fontweight="bold")
        self.figure.tight_layout()
        self.canvas.draw()

    def plot_confidence(self):
        self.figure.clear()
        ax = self.figure.add_subplot(111)
        agg = self._aggregate()
        names = list(agg.keys())
        avg_confs = []
        for n in names:
            if agg[n]["count"] > 0:
                avg_confs.append(agg[n]["total_conf"] / agg[n]["count"])
            else:
                avg_confs.append(0)
        colors = plt.cm.RdYlGn([max(0.3, c) for c in avg_confs])

        bars = ax.barh(names, avg_confs, color=colors, edgecolor="#333", linewidth=0.5)
        for bar, conf in zip(bars, avg_confs):
            ax.text(bar.get_width() + 0.01, bar.get_y() + bar.get_height() / 2,
                    f"{conf:.1%}", ha="left", va="center", fontsize=10)

        ax.set_xlim(0, 1.15)
        ax.set_title("各花朵平均置信度", fontsize=14, fontweight="bold")
        ax.set_xlabel("平均置信度")
        self.figure.tight_layout()
        self.canvas.draw()

    def plot_line(self):
        self.figure.clear()
        ax = self.figure.add_subplot(111)
        confs = []
        for rec in self.records:
            c = rec["置信度"]
            if isinstance(c, (int, float)):
                confs.append(c)
            else:
                confs.append(0)

        ax.plot(range(1, len(confs) + 1), confs, marker="o", markersize=3,
                color="#667eea", linewidth=1.2, alpha=0.8)
        ax.fill_between(range(1, len(confs) + 1), confs, alpha=0.15, color="#667eea")
        ax.set_ylim(0, 1.05)
        ax.set_title("逐张识别置信度变化", fontsize=14, fontweight="bold")
        ax.set_xlabel("图片序号")
        ax.set_ylabel("置信度")
        ax.grid(True, alpha=0.3)
        self.figure.tight_layout()
        self.canvas.draw()


if __name__ == "__main__":
    app = QtWidgets.QApplication(sys.argv)
    window = ImagePlayer()
    window.show()
    sys.exit(app.exec_())
